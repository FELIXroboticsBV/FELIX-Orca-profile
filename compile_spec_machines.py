#!/usr/bin/env python3
import json
import os
import sys

import openpyxl

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else "FELIX_OrcaSlicer_DataSource.xlsx"
OUTPUT_DIR = sys.argv[2] if len(sys.argv) > 2 else "machine"


def fmt_dia(v):
    s = ("%f" % float(v)).rstrip("0")
    return s if not s.endswith(".") else s + "0"


def parse_offsets(raw, count):
    parts = [p.strip() for p in str(raw or "").split(";") if p.strip()]
    if len(parts) != count:
        parts = (parts + ["0x0"] * count)[:count]
    return parts


def build_leaf(row):
    model = row["PrinterModel"]
    is_double = bool(row["IsDouble"])
    n1, n2 = row["Nozzle1_mm"], row["Nozzle2_mm"]
    variant = fmt_dia(row["PrinterVariant"])
    gflavor = row["GcodeFlavor"] or "repetier"

    if is_double:
        nozzle_diameter = [fmt_dia(n1), fmt_dia(n2)]
        extruder_type = ["Direct Drive", "Direct Drive"]
    else:
        nozzle_diameter = [fmt_dia(n1)]
        extruder_type = ["Direct Drive"]

    data = {
        "type": "machine",
        "name": row["Filename"].rsplit(".json", 1)[0],
        "inherits": row["Inherits"],
        "from": "system",
        "instantiation": "true",
        "printer_model": model,
        "printer_variant": variant,
        "gcode_flavor": gflavor,
        "single_extruder_multi_material": "0",
        "extruder_type": extruder_type,
        "nozzle_diameter": nozzle_diameter,

    }

    n = len(nozzle_diameter)
    data["min_layer_height"] = [fmt_dia(row["MinLayer_mm"])] * n
    data["max_layer_height"] = [fmt_dia(row["MaxLayer_mm"])] * n

    if is_double:
        data["extruder_offset"] = parse_offsets(row["ExtruderOffset"], 2)

    return data


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Machines"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    written = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["PrinterModel"]] is None:
            continue
        row = {h: r[idx[h]] for h in headers}
        leaf = build_leaf(row)
        with open(os.path.join(OUTPUT_DIR, row["Filename"]), "w", encoding="utf-8") as f:
            json.dump(leaf, f, indent=2)
            f.write("\n")
        written += 1

    print("Wrote %d machine files to %s" % (written, OUTPUT_DIR))


if __name__ == "__main__":
    main()