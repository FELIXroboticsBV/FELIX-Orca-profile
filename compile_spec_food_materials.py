#!/usr/bin/env python3
import json
import os
import sys

import openpyxl

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else "FELIX_OrcaSlicer_DataSource.xlsx"
OUTPUT_DIR = sys.argv[2] if len(sys.argv) > 2 else "food_material"
SHEET_NAME = sys.argv[3] if len(sys.argv) > 3 else "Food_Material_Data"


def fmt_num(v):
    s = ("%f" % float(v)).rstrip("0")
    return s if not s.endswith(".") else s + "0"


def fmt_ratio_pct(v):
    return fmt_num(float(v) / 100.0)


def parse_list(raw):
    return [p.strip() for p in str(raw or "").split(",") if p.strip()]


def build_leaf(row):
    flow_ratio = fmt_num(row["Flow ratio"])
    retract_dist = fmt_num(row["RetractionDist_mm"])
    retract_extra = fmt_num(row["ExtraRestartDist_mm"])
    retract_speed = fmt_num(row["RetractionSpeed_mms"])
    toolchange_retract = fmt_num(row["ToolchangeRetraction_mm"])
    toolchange_restart = fmt_num(row["ToolchangeRestart_mm"])
    wipe_dist = fmt_num(row["WipeDist_mm"])
    filament_dia = fmt_num(row["FilamentDiameter_mm"])
    vendor = row["Vendor"] or "FELIX Printers"


    data = {
        "type": "filament",
        "name": row["Filename"].rsplit(".json", 1)[0],
        "inherits": row["Inherits"],
        "from": "system",
        "instantiation": "true",
        "compatible_printers": parse_list(row["Compatible printer names"]),
        "filament_diameter": [filament_dia],
        "filament_flow_ratio": [flow_ratio],
        "filament_retraction_length": [retract_dist],
        "filament_retract_restart_extra": [retract_extra],
        "filament_retraction_speed": [retract_speed],
        "filament_deretraction_speed": [retract_speed],
        "filament_retract_length_toolchange": [toolchange_retract],
        "filament_retract_restart_extra_toolchange": [toolchange_restart],
        "filament_wipe_distance": [wipe_dist],
        "filament_vendor": vendor ,
    }

    return data

def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    written = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["Filename"]] is None:
            continue
        row = {h: r[idx[h]] for h in headers}
        leaf = build_leaf(row)
        with open(os.path.join(OUTPUT_DIR, row["Filename"]), "w", encoding="utf-8") as f:
            json.dump(leaf, f, indent=2)
            f.write("\n")
        written += 1

    print("Wrote %d food material files to %s" % (written, OUTPUT_DIR))


if __name__ == "__main__":
    main()