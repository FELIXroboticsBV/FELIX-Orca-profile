#!/usr/bin/env python3
import json
import os
import sys

import openpyxl

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else "FELIX_OrcaSlicer_DataSource.xlsx"
OUTPUT_DIR = sys.argv[2] if len(sys.argv) > 2 else "filament"
SHEET_NAME = sys.argv[3] if len(sys.argv) > 3 else "Plastic_Material_Data"


def fmt_num(v):
    s = ("%f" % float(v)).rstrip("0")
    return s if not s.endswith(".") else s + "0"


def fmt_ratio_pct(v):
    return fmt_num(float(v) / 100.0)


def parse_list(raw):
    return [p.strip() for p in str(raw or "").split(",") if p.strip()]

def fmt_int(v):
    return str(int(round(float(v))))

def is_true(v):
    if isinstance(v, bool):
        return v
    return str(v).strip().upper() == "TRUE"


def build_leaf(row):
    print_temp = fmt_num(row["PrintTemp_C"])
    bed_temp = fmt_num(row["BedTemp_C"])
    chamber_temp = fmt_num(row["ChamberTemp_C"])
    standby_temp = fmt_num(row["StandbyTemp_C"])
    retract_dist = fmt_num(row["RetractionDist_mm"])
    retract_speed = fmt_num(row["RetractionSpeed_mms"])
    retract_extra = fmt_num(row["RetractExtraPrime_mm"])
    flow_ratio = fmt_ratio_pct(row["FlowRatio_pct"])
    max_vol_speed = fmt_num(row["MaxVolumetricSpeed_mm3s"])
    fan_speed = fmt_num(row["FanSpeed_pct"])
    density = fmt_num(row["Density_gcc"])
    vendor = [row["Vendor"]] or "FELIX Printers"

    data = {
        "type": "filament",
        "name": row["Filename"].rsplit(".json", 1)[0],
        "inherits": row["Inherits"],
        "from": "system",
        "instantiation": "true",
        "compatible_printers": parse_list(row["Compatible printer names"]),
        "nozzle_temperature": [print_temp],
        "nozzle_temperature_initial_layer": [print_temp],
        "hot_plate_temp": [bed_temp],
        "hot_plate_temp_initial_layer": [bed_temp],
        "cool_plate_temp": [bed_temp],
        "cool_plate_temp_initial_layer": [bed_temp],
        "eng_plate_temp": [bed_temp],
        "eng_plate_temp_initial_layer": [bed_temp],
        "textured_plate_temp": [bed_temp],
        "textured_plate_temp_initial_layer": [bed_temp],
        "chamber_temperature": [chamber_temp],
        "idle_temperature": [fmt_int(row["StandbyTemp_C"])],
        "filament_retract_length": [retract_dist],
        "filament_retract_speed": [retract_speed],
        "filament_retraction_speed": [retract_speed],
        "filament_retract_restart_extra": [retract_extra],
        "filament_flow_ratio": [flow_ratio],
        "filament_max_volumetric_speed": [max_vol_speed],
        "fan_min_speed": [fan_speed],
        "fan_max_speed": [fan_speed],
        "filament_density": [density],
        "filament_vendor": [vendor] ,
    }

    return data


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    written = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["Filename"]] is None:
            continue
        row = {h: r[idx[h]] for h in headers}
        leaf = build_leaf(row)
        with open(os.path.join(OUTPUT_DIR, row["Filename"]), "w", encoding="utf-8") as f:
            json.dump(leaf, f, indent=2)
            f.write("\n")
        written += 1

    print("Wrote %d filament files to %s" % (written, OUTPUT_DIR))


if __name__ == "__main__":
    main()