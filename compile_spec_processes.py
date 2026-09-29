#!/usr/bin/env python3
import json
import os
import sys

import openpyxl

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else "FELIX_OrcaSlicer_DataSource.xlsx"
OUTPUT_DIR = sys.argv[2] if len(sys.argv) > 2 else "process"
SHEET_NAME = sys.argv[3] if len(sys.argv) > 3 else "Process_Data"


def fmt_num(v):
    s = ("%f" % float(v)).rstrip("0")
    return s if not s.endswith(".") else s + "0"


def fmt_int(v):
    return str(int(v))


def parse_list(raw):
    return [p.strip() for p in str(raw or "").split(",") if p.strip()]


def build_leaf(row):
    layer_height = fmt_num(row["LayerHeight_mm"])
    first_layer_height = fmt_num(row["FirstLayerHeight_mm"])
    line_width = fmt_num(row["LineWidth_mm"])
    shell_layers = fmt_int(row["ShellLayers"])
    support_z = fmt_num(row["SupportZ_mm"])
    infill_speed = fmt_num(row["InfillSpeed_mms"])
    wall_speed = fmt_num(row["WallSpeed_mms"])

    data = {
        "type": "process",
        "name": row["Filename"].rsplit(".json", 1)[0],
        "inherits": row["Inherits"],
        "from": "system",
        "instantiation": "true",
        "compatible_printers": parse_list(row["Compatible printer names"]),
        "layer_height": layer_height,
        "initial_layer_print_height": first_layer_height,
        "line_width": line_width,
        "outer_wall_line_width": line_width,
        "inner_wall_line_width": line_width,
        "top_surface_line_width": line_width,
        "sparse_infill_line_width": line_width,
        "internal_solid_infill_line_width": line_width,
        "initial_layer_line_width": line_width,
        "support_line_width": line_width,
        "top_shell_layers": shell_layers,
        "bottom_shell_layers": shell_layers,
        "support_top_z_distance": support_z,
        "support_bottom_z_distance": support_z,
        "inner_wall_speed": [infill_speed],
        "sparse_infill_speed": [infill_speed],
        "support_speed": [infill_speed],
        "outer_wall_speed": [wall_speed],
        "internal_solid_infill_speed": [wall_speed],
        "top_surface_speed": [wall_speed],
        "initial_layer_speed": [wall_speed],
        "initial_layer_infill_speed": [wall_speed],
    }

    return data


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    written = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["Filename"]] is None:
            continue
        row = {h: r[idx[h]] for h in headers}
        leaf = build_leaf(row)
        with open(os.path.join(OUTPUT_DIR, row["Filename"]), "w", encoding="utf-8") as f:
            json.dump(leaf, f, indent=2)
            f.write("\n")
        written += 1

    print("Wrote %d process files to %s" % (written, OUTPUT_DIR))


if __name__ == "__main__":
    main()